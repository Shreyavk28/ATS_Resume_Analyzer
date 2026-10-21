
import os
import re
import subprocess
import tempfile

import fitz
from docx import Document


def clean_text(text):
    """Normalize extracted resume text."""
    text = re.sub(r"(\w)\n(\w)", r"\1 \2", text)
    text = re.sub(r"\n+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip().lower()


def extract_text_from_pdf(pdf_path):
    """Extract text from a PDF resume."""
    doc = fitz.open(pdf_path)

    try:
        text = ""

        for page in doc:
            text += page.get_text() + "\n"

    finally:
        doc.close()

    return clean_text(text)


def extract_text_from_docx(docx_path):
    """Extract text from a DOCX resume."""
    document = Document(docx_path)

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    # Also extract text from tables, which are common in resumes.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return clean_text("\n".join(paragraphs))


def extract_text_from_doc(doc_path):
    """
    Extract text from a legacy DOC file using Antiword.

    Antiword must be installed and available on the system PATH.
    """
    try:
        result = subprocess.run(
            ["antiword", doc_path],
            capture_output=True,
            text=True,
            check=True,
            timeout=30,
            encoding="utf-8",
            errors="replace",
        )

        return clean_text(result.stdout)

    except FileNotFoundError as exc:
        raise RuntimeError(
            "Legacy DOC support requires Antiword to be installed "
            "on the server. Please convert the file to DOCX or PDF."
        ) from exc

    except subprocess.CalledProcessError as exc:
        raise ValueError(
            "Could not read this DOC file. Please check the file "
            "or convert it to DOCX or PDF."
        ) from exc


def extract_text_from_resume(file_path):
    """Automatically extract text based on the resume extension."""
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension == ".doc":
        return extract_text_from_doc(file_path)

    raise ValueError(
        "Unsupported resume format. Please upload a PDF, DOC, or DOCX file."
    )
